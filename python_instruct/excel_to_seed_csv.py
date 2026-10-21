"""Konverter Excel -> CSV untuk seeder LakmudDataSeeder.

Sumber : Database IPNU IPPNU PAC Kabupaten Malang - MOORA Kehadiran COST Final - MAKESTA ke LAKMUD.xlsx
Keluaran: storage/app/seed/peserta.csv, storage/app/seed/penilaian.csv (ikut repository; keduanya
          memuat data pribadi peserta, lihat docs/deployment-production.md §4.1)

Pemetaan kolom (terverifikasi terhadap sheet sumber):
  DATABASE PESERTA     : header baris 3, data baris 4  -> kolom 0 = No, kolom 5 = Nama
  PENILAIAN 15 MATERI  : header baris 4, data baris 5  -> kolom 0 = No, 5 = No materi, 6 = nama materi, 7..12 = aspek 1..6
Hasil konversi divalidasi ulang sebelum ditulis; skrip berhenti bila ada yang gagal.

Jalankan: python python_instruct/excel_to_seed_csv.py "C:/path/ke/berkas.xlsx"
"""

import csv
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

SHEET_PESERTA = "DATABASE PESERTA"
SHEET_PENILAIAN = "PENILAIAN 15 MATERI"
EXPECTED_PESERTA = 702
EXPECTED_MATERI = 15
COL_PESERTA_NO, COL_PESERTA_NAMA = 0, 5
COL_NILAI_NO, COL_NILAI_MATERI_NO, COL_NILAI_MATERI_NAMA = 0, 5, 6
COL_ASPEK = (7, 8, 9, 10, 11, 12)


def nilai_aspek(row):
    """aspek1 = 0/1 (tidak hadir), aspek2..6 = 10..20 maksimal dua desimal."""
    aspek = []
    for index, kolom in enumerate(COL_ASPEK):
        raw = row[kolom]
        if raw is None or str(raw).strip() == "":
            raise ValueError(f"aspek{index + 1} kosong pada baris No {row[COL_NILAI_NO]}")
        value = float(raw)
        if index == 0:
            if value not in (0, 1):
                raise ValueError(f"aspek1 harus 0/1, dapat {value} (No {row[COL_NILAI_NO]})")
            aspek.append(int(value))
        else:
            if not 10 <= value <= 20:
                raise ValueError(f"aspek{index + 1} di luar 10..20: {value} (No {row[COL_NILAI_NO]})")
            if round(value, 2) != value:
                raise ValueError(f"aspek{index + 1} lebih dari dua desimal: {value} (No {row[COL_NILAI_NO]})")
            aspek.append(value)
    return aspek


def baca_peserta(workbook):
    sheet = workbook[SHEET_PESERTA]
    peserta = {}
    for row in sheet.iter_rows(min_row=4, values_only=True):
        if row is None or row[COL_PESERTA_NO] in (None, ""):
            continue
        no = int(row[COL_PESERTA_NO])
        nama = (row[COL_PESERTA_NAMA] or "").strip()
        if not nama:
            raise ValueError(f"nama kosong pada No {no}")
        if no in peserta:
            raise ValueError(f"No {no} muncul lebih dari sekali di {SHEET_PESERTA}")
        peserta[no] = nama
    return peserta


def baca_penilaian(workbook):
    sheet = workbook[SHEET_PENILAIAN]
    baris = []
    materi = {}
    materi_per_peserta = defaultdict(set)
    for row in sheet.iter_rows(min_row=5, values_only=True):
        if row is None or row[COL_NILAI_NO] in (None, ""):
            continue
        no = int(row[COL_NILAI_NO])
        materi_no = int(row[COL_NILAI_MATERI_NO])
        materi_nama = (row[COL_NILAI_MATERI_NAMA] or "").strip()
        if not materi_nama:
            raise ValueError(f"nama materi kosong pada No {no} materi {materi_no}")
        materi.setdefault(materi_no, materi_nama)
        if materi[materi_no] != materi_nama:
            raise ValueError(f"materi {materi_no} punya dua nama berbeda")
        materi_per_peserta[no].add(materi_no)
        baris.append((no, materi_no, materi_nama, *nilai_aspek(row)))
    return baris, materi, materi_per_peserta


def validasi(peserta, baris, materi, materi_per_peserta):
    masalah = []
    if len(peserta) != EXPECTED_PESERTA:
        masalah.append(f"peserta {len(peserta)} != {EXPECTED_PESERTA}")
    nomor = sorted(peserta)
    if nomor != list(range(1, EXPECTED_PESERTA + 1)):
        masalah.append("nomor peserta tidak berurutan 1..702")
    if len(materi) != EXPECTED_MATERI:
        masalah.append(f"materi {len(materi)} != {EXPECTED_MATERI}")
    if sorted(materi) != list(range(1, EXPECTED_MATERI + 1)):
        masalah.append("nomor materi tidak berurutan 1..15")
    if len(baris) != EXPECTED_PESERTA * EXPECTED_MATERI:
        masalah.append(f"baris penilaian {len(baris)} != {EXPECTED_PESERTA * EXPECTED_MATERI}")
    kurang = [no for no, dipunya in materi_per_peserta.items() if len(dipunya) != EXPECTED_MATERI]
    if kurang:
        masalah.append(f"{len(kurang)} peserta tidak lengkap 15 materi, contoh {kurang[:5]}")
    tanpa_nilai = sorted(set(peserta) - set(materi_per_peserta))
    if tanpa_nilai:
        masalah.append(f"{len(tanpa_nilai)} peserta tanpa baris penilaian, contoh {tanpa_nilai[:5]}")
    if masalah:
        raise SystemExit("VALIDASI GAGAL:\n - " + "\n - ".join(masalah))


def tulis(peserta, baris, target_dir: Path):
    target_dir.mkdir(parents=True, exist_ok=True)
    with (target_dir / "peserta.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["no", "nama"])
        for no in sorted(peserta):
            writer.writerow([no, peserta[no]])
    with (target_dir / "penilaian.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["no", "materi_no", "materi_nama", "aspek1", "aspek2", "aspek3", "aspek4", "aspek5", "aspek6"])
        for no, materi_no, materi_nama, a1, a2, a3, a4, a5, a6 in sorted(baris):
            writer.writerow([no, materi_no, materi_nama, a1, f"{a2:.2f}", f"{a3:.2f}", f"{a4:.2f}", f"{a5:.2f}", f"{a6:.2f}"])


def main() -> int:
    if len(sys.argv) < 2:
        print(__doc__)
        return 2
    sumber = Path(sys.argv[1])
    root = Path(__file__).resolve().parent.parent
    target_dir = root / "storage" / "app" / "seed"

    workbook = openpyxl.load_workbook(sumber, read_only=True, data_only=True)
    peserta = baca_peserta(workbook)
    baris, materi, materi_per_peserta = baca_penilaian(workbook)
    workbook.close()

    validasi(peserta, baris, materi, materi_per_peserta)
    tulis(peserta, baris, target_dir)

    nama_unik = len(set(peserta.values()))
    print(f"peserta          : {len(peserta)} baris ({nama_unik} nama unik, {len(peserta) - nama_unik} nama kembar)")
    print(f"materi           : {len(materi)} (urutan 1..{max(materi)})")
    print(f"baris penilaian  : {len(baris)} ({EXPECTED_MATERI} per peserta, 0 bolong)")
    print(f"aspek1 (0/1)     : {sum(1 for b in baris if b[3] == 1)} baris tidak hadir dari {len(baris)}")
    print(f"tulis            : {target_dir / 'peserta.csv'}")
    print(f"tulis            : {target_dir / 'penilaian.csv'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
